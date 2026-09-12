import os
from pathlib import Path
from langchain_core.documents import Document
import pypdf
import docx

def load_all_docs(folder_path: str) -> list[Document]:
    folder = Path(folder_path)
    if not folder.exists():
        raise FileNotFoundError(f"Document folder does not exist: {folder_path}")
    if not folder.is_dir():
        raise NotADirectoryError(f"Document path is not a folder: {folder_path}")

    try:
        print(f"Loading documents from folder: {folder_path}")
        docs = []

        for path in folder.iterdir():
            if not path.is_file():
                continue

            if path.suffix.lower() == ".pdf":
                docs.extend(read_pdf(str(path)))
            elif path.suffix.lower() == ".docx":
                docs.extend(read_docx(str(path)))
            elif path.suffix.lower() == ".txt":
                docs.extend(read_txt(str(path)))

        return docs
    except Exception as error:
        raise RuntimeError(
            f"Unable to load documents from folder '{folder_path}': {error}"
        ) from error


def read_pdf(file_path: str) -> list[Document]:
    """
    Load a PDF file and return a list of Document objects.

    Args:
        file_path (str): The path to the PDF file.

    """

    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"PDF file does not exist: {file_path}")

    try:
        reader = pypdf.PdfReader(str(path))
        documents = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                documents.append(
                    Document(
                        page_content=text,
                        metadata={"page": i + 1, "FileName": path.name},
                    )
                )

        return documents
    except Exception as error:
        raise RuntimeError(f"Unable to read PDF file '{file_path}'") from error

def read_docx(file_path: str) -> list[Document]:
    """
    Load a DOCX file and return a list of Document objects.

    Args:
        file_path (str): The path to the DOCX file.

    """

    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"DOCX file does not exist: {file_path}")

    try:
        doc = docx.Document(str(path))
        documents = []
        for i, paragraph in enumerate(doc.paragraphs):
            text = paragraph.text
            if text:
                documents.append(
                    Document(
                        page_content=text,
                        metadata={"paragraph": i + 1, "FileName": path.name},
                    )
                )

        return documents
    except Exception as error:
        raise RuntimeError(f"Unable to read DOCX file '{file_path}'") from error


def read_txt(file_path: str) -> list[Document]:
    """
    Load a text file and return it as a Document object.

    Args:
        file_path (str): The path to the text file.

    """

    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"Text file does not exist: {file_path}")

    try:
        with path.open("r", encoding="utf-8") as text_file:
            text = text_file.read()

        if not text:
            return []

        return [Document(page_content=text, metadata={"FileName": path.name})]
    except Exception as error:
        raise RuntimeError(f"Unable to read text file '{file_path}'") from error